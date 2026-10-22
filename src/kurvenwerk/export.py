"""Ausgabe ohne Qt: Tabelle als CSV/xlsx/LaTeX (siunitx), Plot als eigenständiges Python-Skript, Projekt-Export."""
import csv
import re
from pathlib import Path

import numpy as np
import sympy as sp
from sympy.printing.numpy import NumPyPrinter

from . import fit as fitting
from . import plot
from .model import Page, Project, Table
from .rounding import num


def safe_name(s: str) -> str:
    """Datei-/Ordnername ohne unter Windows verbotene Zeichen."""
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", s).strip(" .") or "_"


def rows(t: Table) -> list[list]:
    """Kopfzeile `Größe / Einheit`, dann Werte (None = leere Zelle)."""
    n = max((len(c.values) for c in t.columns), default=0)
    return [[plot.heading(c) for c in t.columns]] + \
        [[c.values[i] if i < len(c.values) else None for c in t.columns] for i in range(n)]


def write_csv(t: Table, path) -> None:
    """Komma als Trenner, Punkt als Dezimalzeichen (RFC 4180) → der Einfüge-Parser liest es zurück."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for r in rows(t):
            w.writerow(["" if v is None else v if isinstance(v, str) else format(v, ".15g") for v in r])


def write_xlsx(tables: list[Table], path) -> None:
    """Ein Blatt je Tabelle."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for t in tables:
        name = re.sub(r"[\[\]:*?/\\]", "_", t.name)[:31] or "Tabelle"
        base, k = name, 2
        while name.lower() in used:
            name = f"{base[:28]}_{k}"
            k += 1
        used.add(name.lower())
        ws = wb.create_sheet(name)
        for r in rows(t):
            ws.append(r)
    wb.save(path)


def _tex(s: str) -> str:
    """Text mit `$…$` bleibt LaTeX, sonst Sonderzeichen maskieren."""
    if "$" in s:
        return s
    m = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
         "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(m.get(ch, ch) for ch in s)


def latex_table(t: Table, rounded: bool = True) -> str:
    """siunitx-S-Spalten; X/Y mit zugehöriger Δ-Spalte werden zu `Wert \\pm Fehler` (DIN-gerundet) zusammengefasst."""
    cols = t.columns
    err, used = {}, set()
    for j, c in enumerate(cols):
        if c.role in ("X", "Y") and (e := plot._error(cols, j)) is not None:
            err[j] = e
            used.add(id(e))
    keep = [j for j, c in enumerate(cols) if id(c) not in used]
    n = max((len(c.values) for c in cols), default=0)
    get = lambda c, i: c.values[i] if i < len(c.values) else None  # noqa: E731
    head = []
    for j in keep:
        c = cols[j]
        head.append("{" + _tex(c.name) + (rf" / \unit{{{c.unit}}}" if c.unit else "") + "}")
    body = []
    for i in range(n):
        cells = []
        for j in keep:
            v = get(cols[j], i)
            e = get(err[j], i) if j in err else None
            cells.append("{}" if v is None else num(v, abs(e) if e else None, rounded, tex=True))
        body.append("    " + " & ".join(cells) + r" \\")
    return "\n".join([r"% benötigt \usepackage{siunitx} (v3) und \usepackage{booktabs}",
                      rf"\begin{{tabular}}{{{'S' * len(keep)}}}", r"    \toprule",
                      "    " + " & ".join(head) + r" \\", r"    \midrule", *body, r"    \bottomrule",
                      r"\end{tabular}", ""])


def _arr(a) -> str:
    return f"np.array({[float(v) for v in a]!r})"


def _fit_code(s: plot.Series, i: int) -> tuple[list[str], tuple] | None:
    """def fit_i(x, a=…, …) → numpy-Code des Modells mit den gefitteten Werten; None, wenn der Fit scheitert."""
    try:
        r, m = plot.run_fit(s)
    except ValueError:
        return None
    expr, names = fitting.model(s.ycol.fit["formula"])
    safe = [n if n.isidentifier() and n not in ("x", "np", "numpy") else f"p{k}" for k, n in enumerate(names)]
    expr = expr.subs({sp.Symbol(n): sp.Symbol(k) for n, k in zip(names, safe)})
    args = ", ".join(f"{k}={float(v)!r}" for k, v in zip(safe, r.values))
    return [f"def fit_{i}(x, {args}):", f"    return {NumPyPrinter().doprint(expr)}"], (s.x[m].min(), s.x[m].max())


def python_script(page: Page, style: dict) -> str:
    """Eigenständiges Skript (numpy + matplotlib), das den Plot der Seite nachbaut.
    ponytail: ohne Konfidenzband und Residuenplot, nachrüsten wenn gefragt."""
    st, pl = {**plot.STYLE, **style}, {**plot.PLOT, **page.plot}
    colors = [c.strip() for c in str(st["colors"]).split(",") if c.strip()] or ["C0"]
    swap = pl["swap"]
    all_s = [s for t in page.tables for s in plot.series(t)]
    L = [f'"""Plot „{page.title}“ — erzeugt von Kurvenwerk."""', "import matplotlib.pyplot as plt",
         "import numpy", "import numpy as np", "",
         f'plt.rcParams.update({{"font.size": {st["fontsize"]!r}, "font.family": {st["font"]!r}}})',
         f"fig, ax = plt.subplots(figsize=({st['export_w']!r} / 2.54, {st['export_h']!r} / 2.54), "
         'layout="constrained")',
         f"ax.grid({bool(st['grid'])}, alpha=0.3)"]
    has2 = any({**plot.SERIES, **s.ycol.style}["axis"] == "y2" for s in all_s)
    if has2:
        L.append(f"ax2 = ax.{'twiny' if swap else 'twinx'}()")
    names = {"x": [], "y": [], "y2": []}
    lw = st["linewidth"]
    for i, s in enumerate(all_s):
        o = {**plot.SERIES, **s.ycol.style}
        ms = o["markersize"] or st["markersize"]
        a = "ax2" if o["axis"] == "y2" and has2 else "ax"
        x, y, dx, dy = (s.y, s.x, s.dy, s.dx) if swap else (s.x, s.y, s.dx, s.dy)
        if not o["errorbars"]:
            dx = dy = None
        ls = o["linestyle"] or st["linestyle"]
        color, marker = o["color"] or colors[i % len(colors)], o["marker"] or st["marker"]
        out = np.isin(s.rows, s.ycol.fit.get("excluded", []))
        L += ["", f"# {plot.heading(s.ycol)} über {plot.heading(s.xcol)}", f"x, y = {_arr(x[~out])}, {_arr(y[~out])}"]
        L.append(f"dx = {_arr(dx[~out]) if dx is not None else None}")
        L.append(f"dy = {_arr(dy[~out]) if dy is not None else None}")
        L.append(f"{a}.errorbar(x, y, xerr=dx, yerr=dy, label={plot.safe(o['label'] or plot.heading(s.ycol))!r}, "
                 f"color={color!r}, marker={marker!r}, markersize={ms!r}, "
                 f"linestyle={'none' if ls == 'none' else ls!r}, linewidth={lw!r}, capsize={st['capsize']!r}, "
                 f"elinewidth={lw * 0.8!r})")
        if out.any():
            L.append(f"{a}.plot({_arr(x[out])}, {_arr(y[out])}, linestyle='none', marker={marker!r}, "
                     f"markersize={ms!r}, markerfacecolor='none', markeredgecolor={color!r}, alpha=0.6)"
                     "  # vom Fit ausgeschlossen")
        if s.ycol.fit.get("formula") and (fc := _fit_code(s, i)):
            code, (lo, hi) = fc
            space = "geomspace" if pl["xlog"] and lo > 0 else "linspace"
            pts = f"fit_{i}(xs), xs" if swap else f"xs, fit_{i}(xs)"
            L += code + [f"xs = np.{space}({float(lo)!r}, {float(hi)!r}, 400)",
                         f"{a}.plot({pts}, color={color!r}, linewidth={lw * 1.2!r}, alpha={plot.fit_alpha(s, st)!r}, label={'Fit ' + s.ycol.name!r})"]
        names["x"].append(plot.heading(s.xcol))
        names["y2" if a == "ax2" else "y"].append(plot.heading(s.ycol))
    hx, vy = ("y", "x") if swap else ("x", "y")
    L.append("")
    for a, axis, lab, log, lo, hi in [("ax", hx, pl["xlabel"] or plot._auto(names["x"]), pl["xlog"], pl["xmin"],
                                       pl["xmax"]),
                                      ("ax", vy, pl["ylabel"] or plot._auto(names["y"]), pl["ylog"], pl["ymin"],
                                       pl["ymax"]),
                                      ("ax2", vy, pl["y2label"] or plot._auto(names["y2"]), pl["y2log"], pl["y2min"],
                                       pl["y2max"])]:
        if a == "ax2" and not has2:
            continue
        L.append(f"{a}.set_{axis}label({plot.safe(lab)!r})")
        if log:
            L.append(f'{a}.set_{axis}scale("log")')
        if lo is not None or hi is not None:
            L.append(f"{a}.set_{axis}lim({lo!r}, {hi!r})")
    if pl["title"]:
        L.append(f"ax.set_title({plot.safe(pl['title'])!r})")
    if pl["legend"] and all_s:
        L.append("ax.legend(handles=[h for a in fig.axes for h in a.get_legend_handles_labels()[0]])")
    L += ["", 'if __name__ == "__main__":', "    plt.show()", ""]
    return "\n".join(L)


def export_project(project: Project, name: str, dest) -> Path:
    """<dest>/<name>_Export/ mit projekt.json und je Seite 01_<Titel>/ mit plot.png + <Tabelle>.csv.
    Vorhandene Dateien werden überschrieben, fremde bleiben liegen."""
    root = Path(dest) / f"{safe_name(name)}_Export"
    root.mkdir(parents=True, exist_ok=True)
    project.save(root / "projekt.json")
    st = {**plot.STYLE, **project.style}
    for i, p in enumerate(project.pages, 1):
        d = root / f"{i:02d}_{safe_name(p.title)}"
        d.mkdir(exist_ok=True)
        plot.export(p, project.style, str(d / "plot.png"), st["export_w"], st["export_h"], st["export_dpi"])
        used = set()
        for t in p.tables:
            n, k = safe_name(t.name), 2
            while n.lower() in used:
                n = f"{safe_name(t.name)}_{k}"
                k += 1
            used.add(n.lower())
            write_csv(t, d / f"{n}.csv")
    return root
