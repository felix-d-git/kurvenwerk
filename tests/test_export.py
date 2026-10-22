import csv
import runpy

import matplotlib
import numpy as np
import pytest

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

from kurvenwerk import export  # noqa: E402
from kurvenwerk.model import Column, Page, Project, Table  # noqa: E402

X = [1.0, 2.0, 3.0, 4.0, 5.0]
Y = [2.1, 3.9, 6.2, 7.8, 10.1]  # synthetisch, ≈ 2x


def page(fit=None, **plot):
    t = Table("Mess:1", [Column("t", "s", "X", list(X)),
                         Column("U", "V", "Y", list(Y), fit=fit or {}),
                         Column("ΔU", "V", "Δy", [0.1, 0.1, 0.2, 0.2, 0.3])])
    return Page("Seite/1", [t], plot=plot)


def test_csv_roundtrip(tmp_path):
    t = page().tables[0]
    t.columns[1].values[2] = None
    export.write_csv(t, tmp_path / "a.csv")
    with open(tmp_path / "a.csv", encoding="utf-8", newline="") as f:
        r = list(csv.reader(f))
    assert r[0] == ["t / s", "U / V", "ΔU / V"]
    assert r[3] == ["3", "", "0.2"]
    assert float(r[5][1]) == 10.1


def test_xlsx(tmp_path):
    from openpyxl import load_workbook
    p = page()
    export.write_xlsx(p.tables + [Table("Mess?1", [Column("a", "", "X", [1.5])])], tmp_path / "a.xlsx")
    wb = load_workbook(tmp_path / "a.xlsx")
    assert wb.sheetnames == ["Mess_1", "Mess_1_2"]
    assert wb["Mess_1"]["B6"].value == 10.1


def test_latex():
    s = export.latex_table(page().tables[0])
    assert r"\begin{tabular}{SS}" in s  # ΔU in U zusammengefasst
    assert r"{t / \unit{s}} & {U / \unit{V}} \\" in s
    assert r"1 & 2.10 \pm 0.10 \\" in s
    assert r"5 & 10.1 \pm 0.3 \\" in s
    assert export.latex_table(Table("x", [Column("a_b%", "", "Y", [1.0, None])])).count("{}") == 1


def _run(src, tmp_path):
    f = tmp_path / "plot_script.py"
    f.write_text(src, encoding="utf-8")
    plt.close("all")
    g = runpy.run_path(str(f))  # __name__ != "__main__" → kein plt.show()
    return g["fig"]


def test_script_reproduces_data_and_fit(tmp_path):
    p = page({"formula": "a*x + b", "excluded": [4]}, title="T", ylog=True)
    fig = _run(export.python_script(p, {}), tmp_path)
    ax = fig.axes[0]
    lines = [ln for ln in ax.get_lines() if len(ln.get_xdata())]
    data, fitline = lines[0], lines[-1]
    hollow = next(ln for ln in lines if ln.get_markerfacecolor() == "none")
    assert np.allclose(data.get_ydata(), Y[:4]) and np.allclose(hollow.get_xdata(), [5.0])
    # Fit mit Gewichten 1/ΔU² über die ersten 4 Punkte, Kurve über 1…4
    xs, ys = fitline.get_xdata(), fitline.get_ydata()
    assert xs[0] == 1 and xs[-1] == 4
    a, b = np.polyfit(X[:4], Y[:4], 1, w=1 / np.array([0.1, 0.1, 0.2, 0.2]))
    assert np.allclose(ys, a * xs + b, rtol=1e-6)
    assert ax.get_yscale() == "log" and ax.get_title() == "T" and ax.get_xlabel() == "t / s"


def test_script_odd_param_names(tmp_path):
    p = page({"formula": "A_{0}*exp(x/tau) + lambda"})
    p.tables[0].columns[1].values = [1 + np.exp(v / 3) for v in X]
    src = export.python_script(p, {})
    compile(src, "s", "exec")
    _run(src, tmp_path)


@pytest.mark.filterwarnings("ignore")
def test_project_export(tmp_path):
    pr = Project([page(), page()], {})
    pr.pages[1].title = "Seite/1"  # gleicher Titel → Nummer trennt
    root = export.export_project(pr, "Ver:such", tmp_path)
    assert root.name == "Ver_such_Export"
    assert (root / "projekt.json").exists()
    for d in ("01_Seite_1", "02_Seite_1"):
        assert (root / d / "plot.png").stat().st_size > 1000
        assert (root / d / "Mess_1.csv").exists()
    Project.load(root / "projekt.json")


def test_series_style_no_errorbars_small_markers(tmp_path):
    from matplotlib.figure import Figure
    from kurvenwerk import plot
    p = page()
    p.tables[0].columns[1].style = {"errorbars": False, "marker": ".", "markersize": 2.0, "linestyle": "-"}
    fig = Figure()
    plot.draw(fig, p, {})
    ax = fig.axes[0]
    assert not ax.collections  # keine Fehlerbalken-LineCollection
    ln = ax.get_lines()[0]
    assert ln.get_marker() == "." and ln.get_markersize() == 2.0 and ln.get_linestyle() == "-"
    src = export.python_script(p, {})
    assert "dy = None" in src
    ln = _run(src, tmp_path).axes[0].get_lines()[0]
    assert ln.get_marker() == "." and ln.get_markersize() == 2.0


def test_fit_alpha_and_default_colors(tmp_path):
    from matplotlib.colors import to_hex
    from matplotlib.figure import Figure
    from kurvenwerk import plot
    p = page({"formula": "a*x + b"})
    fig = Figure()
    plot.draw(fig, p, {"fit_alpha": 0.4})
    lines = fig.axes[0].get_lines()
    assert to_hex(lines[0].get_color()) == "#1f77b4"  # Reihe 1 blau
    assert lines[-1].get_alpha() == 0.4 and to_hex(lines[-1].get_color()) == "#1f77b4"
    p.tables[0].columns[1].style = {"fit_alpha": 7, "color": "red"}  # Reihe schlägt Projekt, begrenzt auf 1
    fig = Figure()
    plot.draw(fig, p, {"fit_alpha": 0.4})
    assert fig.axes[0].get_lines()[-1].get_alpha() == 1.0
    assert _run(export.python_script(p, {"fit_alpha": 0.4}), tmp_path).axes[0].get_lines()[-1].get_alpha() == 1.0
